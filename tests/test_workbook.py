from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import load_workbook

from hitam_cli.portal import AnswerSheet
from hitam_cli.workbook import META_SHEET, read_marks_workbook, write_workbook


class WorkbookTests(TestCase):
    def test_workbook_contains_clickable_links_and_sheet_numbers_as_text(self):
        rows = [
            AnswerSheet("000001", "https://hitamexams.org/ScriptPDF/CS101/000001.pdf"),
            AnswerSheet("123456", "https://hitamexams.org/ScriptPDF/CS101/123456.pdf"),
        ]
        with TemporaryDirectory() as directory:
            output = Path(directory) / "answer-sheets.xlsx"
            saved = write_workbook(rows, output, subject_code="CS101", bundle_no="B17")
            workbook = load_workbook(saved)
            sheet = workbook["Answer sheets"]

            self.assertEqual(sheet["A1"].value, "Sheet No")
            self.assertEqual(sheet["B1"].value, "PDF Link")
            self.assertEqual(sheet["A2"].value, "000001")
            self.assertEqual(sheet["A2"].number_format, "@")
            self.assertEqual(sheet["B2"].hyperlink.target, rows[0].pdf_url)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:B3")
            self.assertEqual(workbook[META_SHEET]["B1"].value, "2")
            self.assertEqual(workbook[META_SHEET]["B2"].value, "CS101")
            self.assertEqual(workbook[META_SHEET]["B3"].value, "B17")
            self.assertEqual(workbook[META_SHEET].sheet_state, "veryHidden")
            workbook.close()

    def test_existing_output_is_not_overwritten(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "existing.xlsx"
            output.write_text("keep me", encoding="utf-8")
            with self.assertRaises(FileExistsError):
                write_workbook(
                    [AnswerSheet("123456", "https://example.test/123456.pdf")],
                    output,
                    subject_code="CS101",
                    bundle_no="B17",
                )
            self.assertEqual(output.read_text(encoding="utf-8"), "keep me")

    def test_empty_workbook_is_rejected(self):
        with TemporaryDirectory() as directory, self.assertRaises(ValueError):
            write_workbook(
                [],
                Path(directory) / "empty.xlsx",
                subject_code="CS101",
                bundle_no="B17",
            )

    def test_completed_marks_workbook_accepts_decimal_question_values(self):
        rows = [
            AnswerSheet("000001", "https://hitamexams.org/ScriptPDF/CS101/000001.pdf"),
            AnswerSheet("123456", "https://hitamexams.org/ScriptPDF/CS101/123456.pdf"),
        ]
        with TemporaryDirectory() as directory:
            output = Path(directory) / "marks.xlsx"
            write_workbook(rows, output, subject_code="CS101", bundle_no="B17")
            workbook = load_workbook(output)
            sheet = workbook["Answer sheets"]
            sheet["C1"] = "Q1"
            sheet["D1"] = "Q2(a)"
            sheet["C2"] = 4
            sheet["D2"] = 2.5
            sheet["C3"] = 3
            sheet["D3"] = 1
            workbook.save(output)
            workbook.close()

            result = read_marks_workbook(output)

            self.assertEqual(result.subject_code, "CS101")
            self.assertEqual(result.bundle_no, "B17")
            self.assertEqual(result.question_headers, ("Q1", "Q2(a)"))
            self.assertEqual(str(result.students[0].marks["1"]), "4")
            self.assertEqual(str(result.students[0].marks["2(a)"]), "2.5")

    def test_marks_workbook_rejects_blank_and_formula_marks(self):
        rows = [
            AnswerSheet("000001", "https://hitamexams.org/ScriptPDF/CS101/000001.pdf")
        ]
        with TemporaryDirectory() as directory:
            output = Path(directory) / "marks.xlsx"
            write_workbook(rows, output, subject_code="CS101", bundle_no="B17")
            workbook = load_workbook(output)
            workbook["Answer sheets"]["C1"] = "Q1"
            workbook.save(output)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "mark is blank"):
                read_marks_workbook(output)

            workbook = load_workbook(output)
            workbook["Answer sheets"]["C2"] = "=1+1"
            workbook.save(output)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "formulas are not allowed"):
                read_marks_workbook(output)
