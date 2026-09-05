from __future__ import annotations

import os
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, skipUnless

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from hitam_cli.portal import (
    AnswerSheet,
    EvaluationInputs,
    PortalError,
    QuestionSpec,
    StudentSchema,
    _save_student,
    extract_from_page,
    upload_from_page,
)
from hitam_cli.workbook import read_marks_workbook, write_workbook


@skipUnless(
    os.getenv("HITAM_RUN_BROWSER_TESTS") == "1", "browser integration test is opt-in"
)
class MockBrowserFlowTests(TestCase):
    def test_mocked_portal_flow_captures_actual_pdf_requests(self):
        html = """
        <button id="faculty">FACULTY LOGIN</button>
        <div id="stage"></div>
        <script>
        const stage = document.querySelector('#stage');
        const field = label => `<label>${label}<input aria-label="${label}"></label><button>VALIDATE</button>`;
        document.querySelector('#faculty').onclick = () => {
          stage.innerHTML = '<select><option>B.Tech</option></select><button>VALIDATE</button>';
          stage.querySelector('button').onclick = () => { stage.innerHTML = field('USER CODE'); wire('BUNDLE NO'); };
        };
        function wire(next) {
          stage.querySelector('button').onclick = () => {
            if (next === 'BUNDLE NO') { stage.innerHTML = field(next); wire('BUNDLE KEY'); }
            else if (next === 'BUNDLE KEY') { stage.innerHTML = field(next); wire('SHEETS'); }
            else {
              stage.innerHTML = '<button>000001</button><button>123456</button><button>123456</button>';
              for (const button of stage.querySelectorAll('button')) {
                button.onclick = () => fetch(`https://mock.local/ScriptPDF/CS101/${button.textContent}.pdf`);
              }
            }
          };
        }
        </script>
        """
        inputs = EvaluationInputs("USER", "CS101", "BUNDLE", "KEY")
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.route(
                    "https://hitamexams.org/MidEvaluation/",
                    lambda route: route.fulfill(body=html),
                )
                page.route(
                    "https://mock.local/**",
                    lambda route: route.fulfill(body="%PDF-mock"),
                )
                rows = extract_from_page(page, inputs, timeout_ms=5_000)
            finally:
                browser.close()

        self.assertEqual([row.sheet_no for row in rows], ["000001", "123456"])
        self.assertTrue(rows[0].pdf_url.endswith("/CS101/000001.pdf"))

    def test_mocked_upload_preflights_then_overwrites_every_student(self):
        html = """
        <button id="faculty">FACULTY LOGIN</button>
        <div id="stage"></div>
        <script>
        const stage = document.querySelector('#stage');
        window.saved = [];
        const field = label => `<label>${label}<input aria-label="${label}"></label><button>VALIDATE</button>`;
        document.querySelector('#faculty').onclick = () => {
          stage.innerHTML = '<select><option>B.Tech</option></select><button>VALIDATE</button>';
          stage.querySelector('button').onclick = () => { stage.innerHTML = field('USER CODE'); wire('BUNDLE NO'); };
        };
        function wire(next) {
          stage.querySelector('button').onclick = () => {
            if (next === 'BUNDLE NO') { stage.innerHTML = field(next); wire('BUNDLE KEY'); }
            else if (next === 'BUNDLE KEY') { stage.innerHTML = field(next); wire('SHEETS'); }
            else { showSheets(); }
          };
        }
        function showSheets() {
          stage.innerHTML = '<div id="buttons"><button>000001</button><button>123456</button></div><div id="marks"></div>';
          for (const button of document.querySelectorAll('#buttons button')) {
            button.onclick = () => showStudent(button.textContent);
          }
        }
        function showStudent(sheet) {
          fetch(`https://mock.local/ScriptPDF/CS101/${sheet}.pdf`);
          fetch('https://mock.local/MIDScrutinyInternalMarksEntry', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({StudId: sheet, SectionId: 0})
          });
          document.querySelector('#marks').innerHTML = `
            <table><thead><tr><th><label>Script Code: ${sheet}</label></th></tr>
            <tr><th>Q.NO</th><th>MAX MARKS</th><th>MARKS</th></tr></thead>
            <tbody><tr><td>1</td><td>5</td><td><input value="1"></td>
            <tr><td>2(a)</td><td>3</td><td><input value="1"></td></tbody></table>
            <button id="save">Save</button>`;
          const inputs = [...document.querySelectorAll('#marks input')];
          for (const input of inputs) {
            input.onblur = () => fetch(
              'https://mock.local/MIDValuaterSavingMarksCheckTotalGridNew',
              {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({Total: inputs.reduce((sum, item) => sum + Number(item.value || 0), 0)})
              }
            );
          }
          document.querySelector('#save').onclick = () => {
            window.saved.push({sheet, marks: inputs.map(input => input.value)});
            fetch('https://mock.local/MIDSecuritySavingMarksNew', {method: 'POST'});
          };
        }
        </script>
        """
        requests: list[str] = []
        with TemporaryDirectory() as directory:
            path = Path(directory) / "marks.xlsx"
            write_workbook(
                [
                    AnswerSheet(
                        "000001", "https://mock.local/ScriptPDF/CS101/000001.pdf"
                    ),
                    AnswerSheet(
                        "123456", "https://mock.local/ScriptPDF/CS101/123456.pdf"
                    ),
                ],
                path,
                subject_code="CS101",
                bundle_no="B17",
            )
            workbook = load_workbook(path)
            sheet = workbook["Answer sheets"]
            sheet["C1"], sheet["D1"] = "Q1", "Q2(a)"
            sheet["C2"], sheet["D2"] = 4, 2.5
            sheet["C3"], sheet["D3"] = 3, 2
            workbook.save(path)
            workbook.close()
            marks_workbook = read_marks_workbook(path)

            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                try:
                    page = browser.new_page()
                    page.route(
                        "https://hitamexams.org/MidEvaluation/",
                        lambda route: route.fulfill(body=html),
                    )

                    def respond(route):
                        requests.append(route.request.url)
                        if "MIDScrutinyInternalMarksEntry" in route.request.url:
                            route.fulfill(
                                json={
                                    "MarksDisplayTable1": [
                                        {"QNo": "1", "QMaxMarks": 5, "Marks": 1},
                                        {
                                            "QNo": "2(a)",
                                            "QMaxMarks": 3,
                                            "Marks": 1,
                                        },
                                    ]
                                }
                            )
                        elif "MIDSecuritySavingMarksNew" in route.request.url:
                            route.fulfill(
                                json={"SavingMarksToGrid": "SuccessFully Saved Data."}
                            )
                        elif "CheckTotal" in route.request.url:
                            route.fulfill(json=route.request.post_data_json)
                        else:
                            route.fulfill(body="%PDF-mock")

                    page.route("https://mock.local/**", respond)
                    summaries = []
                    saved = upload_from_page(
                        page,
                        EvaluationInputs("USER", "CS101", "B17", "KEY"),
                        marks_workbook,
                        lambda summary: summaries.append(summary) or True,
                        timeout_ms=5_000,
                    )
                    portal_saved = page.evaluate("window.saved")
                finally:
                    browser.close()

        self.assertEqual(saved, 2)
        self.assertEqual(summaries[0].mark_count, 4)
        self.assertEqual(
            portal_saved,
            [
                {"sheet": "000001", "marks": ["4", "2.5"]},
                {"sheet": "123456", "marks": ["3", "2"]},
            ],
        )
        self.assertFalse(any("Finalization" in url for url in requests))

    def test_blank_student_waits_for_delayed_table_before_filling_all_questions(self):
        html = """
        <button id="faculty">FACULTY LOGIN</button>
        <div id="stage"></div>
        <script>
        const stage = document.querySelector('#stage');
        window.saved = [];
        const field = label => `<label>${label}<input aria-label="${label}"></label><button>VALIDATE</button>`;
        document.querySelector('#faculty').onclick = () => {
          stage.innerHTML = '<select><option>B.Tech</option></select><button>VALIDATE</button>';
          stage.querySelector('button').onclick = () => { stage.innerHTML = field('USER CODE'); wire('BUNDLE NO'); };
        };
        function wire(next) {
          stage.querySelector('button').onclick = () => {
            if (next === 'BUNDLE NO') { stage.innerHTML = field(next); wire('BUNDLE KEY'); }
            else if (next === 'BUNDLE KEY') { stage.innerHTML = field(next); wire('SHEETS'); }
            else { showSheets(); }
          };
        }
        function questionRows(rows) {
          return rows.map(row => `<tr><td>${row.QNo}</td><td>${row.QMaxMarks}</td><td><input value="${row.Marks}"></td></tr>`).join('');
        }
        function renderStudent(sheet, rows) {
          document.querySelector('#marks').innerHTML = `
            <table><thead><tr><th><label>Script Code: ${sheet}</label></th></tr>
            <tr><th>Q.NO</th><th>MAX MARKS</th><th>MARKS</th></tr></thead>
            <tbody>${questionRows(rows)}</tbody></table><button id="save">Save</button>`;
          const inputs = [...document.querySelectorAll('#marks tbody input')];
          for (const input of inputs) {
            input.onblur = () => setTimeout(() => fetch(
              'https://mock.local/MIDValuaterSavingMarksCheckTotalGridNew',
              {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({Total: inputs.reduce((sum, item) => sum + Number(item.value || 0), 0)})
              }
            ), 15);
          }
          document.querySelector('#save').onclick = () => {
            window.saved.push({sheet, marks: inputs.map(input => input.value)});
            fetch('https://mock.local/MIDSecuritySavingMarksNew', {method: 'POST'});
          };
        }
        function showSheets() {
          stage.innerHTML = '<div id="buttons"><button>000001</button><button>000002</button></div><div id="marks"></div>';
          const stale = Array.from({length: 11}, (_, index) => ({
            QNo: String(index + 1), QMaxMarks: index < 5 ? 2 : 5, Marks: ''
          }));
          renderStudent('999999', stale);
          for (const button of document.querySelectorAll('#buttons button')) {
            button.onclick = () => showStudent(button.textContent);
          }
        }
        function showStudent(sheet) {
          fetch(`https://mock.local/ScriptPDF/CS101/${sheet}.pdf`);
          fetch('https://mock.local/MIDScrutinyInternalMarksEntry', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({StudId: sheet, SectionId: 0})
          }).then(response => response.json()).then(payload => {
            setTimeout(() => renderStudent(sheet, payload.MarksDisplayTable1), 120);
          });
        }
        </script>
        """
        first_marks = [0, 0, 2, 0.5, 1, 0, 3.5, 0, 2.5, 4.5, 0]
        second_marks = [2, 1.5, 1, 2, 0, 5, 4.5, 0, 3, 2.5, 1]
        maxima = [2, 2, 2, 2, 2, 5, 5, 5, 5, 5, 5]

        with TemporaryDirectory() as directory:
            path = Path(directory) / "blank-students.xlsx"
            write_workbook(
                [
                    AnswerSheet(
                        "000001", "https://mock.local/ScriptPDF/CS101/000001.pdf"
                    ),
                    AnswerSheet(
                        "000002", "https://mock.local/ScriptPDF/CS101/000002.pdf"
                    ),
                ],
                path,
                subject_code="CS101",
                bundle_no="B17",
            )
            workbook = load_workbook(path)
            sheet = workbook["Answer sheets"]
            for column, question in enumerate(range(1, 12), start=3):
                sheet.cell(1, column, f"Q{question}")
                sheet.cell(2, column, first_marks[question - 1])
                sheet.cell(3, column, second_marks[question - 1])
            workbook.save(path)
            workbook.close()
            marks_workbook = read_marks_workbook(path)

            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                try:
                    page = browser.new_page()
                    page.route(
                        "https://hitamexams.org/MidEvaluation/",
                        lambda route: route.fulfill(body=html),
                    )

                    def respond(route):
                        url = route.request.url
                        if "MIDScrutinyInternalMarksEntry" in url:
                            route.fulfill(
                                json={
                                    "MarksDisplayTable1": [
                                        {
                                            "QNo": str(index + 1),
                                            "QMaxMarks": maximum,
                                            "Marks": "",
                                        }
                                        for index, maximum in enumerate(maxima)
                                    ]
                                }
                            )
                        elif "MIDSecuritySavingMarksNew" in url:
                            route.fulfill(
                                json={"SavingMarksToGrid": "SuccessFully Saved Data."}
                            )
                        elif "CheckTotal" in url:
                            route.fulfill(json=route.request.post_data_json)
                        else:
                            route.fulfill(body="%PDF-mock")

                    page.route("https://mock.local/**", respond)
                    saved = upload_from_page(
                        page,
                        EvaluationInputs("USER", "CS101", "B17", "KEY"),
                        marks_workbook,
                        lambda summary: True,
                        timeout_ms=5_000,
                    )
                    portal_saved = page.evaluate("window.saved")
                finally:
                    browser.close()

        self.assertEqual(saved, 2)
        self.assertEqual(
            portal_saved,
            [
                {
                    "sheet": "000001",
                    "marks": [str(mark) for mark in first_marks],
                },
                {
                    "sheet": "000002",
                    "marks": [str(mark) for mark in second_marks],
                },
            ],
        )

    def test_changed_input_is_detected_before_save(self):
        html = """
        <button id="student">000001</button>
        <div id="marks"></div>
        <script>
        window.saveClicks = 0;
        const rows = [
          {QNo: '1', QMaxMarks: 5, Marks: ''},
          {QNo: '2', QMaxMarks: 5, Marks: ''}
        ];
        function renderStudent() {
          document.querySelector('#marks').innerHTML = `
            <table><thead><tr><th><label>Script Code: 000001</label></th></tr>
            <tr><th>Q.NO</th><th>MAX MARKS</th><th>MARKS</th></tr></thead>
            <tbody><tr><td>1</td><td>5</td><td><input data-question="1" value=""></td></tr>
            <tr><td>2</td><td>5</td><td><input data-question="2" value=""></td></tr></tbody></table>
            <button id="save">Save</button>`;
          const inputs = [...document.querySelectorAll('#marks input')];
          for (const input of inputs) {
            input.onblur = () => {
              const total = inputs.reduce((sum, item) => sum + Number(item.value || 0), 0);
              if (input.dataset.question === '2') inputs[0].value = '';
              fetch('https://mock.local/MIDValuaterSavingMarksCheckTotalGridNew', {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({Total: total})
              });
            };
          }
          document.querySelector('#save').onclick = () => window.saveClicks++;
        }
        document.querySelector('#student').onclick = () => {
          fetch('https://mock.local/MIDScrutinyInternalMarksEntry', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({StudId: 'student-1', SectionId: 0})
          });
          renderStudent();
        };
        </script>
        """
        schema = StudentSchema(
            "000001",
            "https://mock.local/ScriptPDF/CS101/000001.pdf",
            (
                QuestionSpec("1", "1", Decimal("5")),
                QuestionSpec("2", "2", Decimal("5")),
            ),
        )

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.set_content(html)

                def respond(route):
                    if "MIDScrutinyInternalMarksEntry" in route.request.url:
                        route.fulfill(json={"MarksDisplayTable1": rows})
                    else:
                        route.fulfill(json=route.request.post_data_json)

                rows = [
                    {"QNo": "1", "QMaxMarks": 5, "Marks": ""},
                    {"QNo": "2", "QMaxMarks": 5, "Marks": ""},
                ]
                page.route("https://mock.local/**", respond)
                with self.assertRaisesRegex(
                    PortalError, "question 1: the portal input is blank"
                ):
                    _save_student(
                        page,
                        schema,
                        {"1": Decimal("4"), "2": Decimal("3")},
                        5_000,
                    )
                save_clicks = page.evaluate("window.saveClicks")
            finally:
                browser.close()

        self.assertEqual(save_clicks, 0)
