"""Excel workbook generation for captured answer-sheet links."""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill

from .portal import AnswerSheet, normalise_question_label

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MARKS_SHEET = "Answer sheets"
META_SHEET = "__hitam_meta"
WORKBOOK_FORMAT_VERSION = "2"


@dataclass(frozen=True)
class StudentMarks:
    sheet_no: str
    pdf_url: str
    marks: dict[str, Decimal]


@dataclass(frozen=True)
class MarksWorkbook:
    path: Path
    subject_code: str
    bundle_no: str
    expected_sheet_count: int
    question_headers: tuple[str, ...]
    students: tuple[StudentMarks, ...]


def _write_metadata(
    workbook: Workbook,
    *,
    subject_code: str,
    bundle_no: str,
    sheet_count: int,
    exported_at: datetime,
) -> None:
    metadata = workbook.create_sheet(META_SHEET)
    metadata.append(["HITAM Workbook Format", WORKBOOK_FORMAT_VERSION])
    metadata.append(["Subject Code", subject_code])
    metadata.append(["Bundle No", bundle_no])
    metadata.append(["Exported At", exported_at.isoformat(timespec="seconds")])
    metadata.append(["Expected Sheet Count", sheet_count])
    metadata.sheet_state = "veryHidden"


def write_workbook(
    answer_sheets: list[AnswerSheet],
    output_path: Path,
    *,
    subject_code: str,
    bundle_no: str,
    exported_at: datetime | None = None,
) -> Path:
    """Atomically write a formatted workbook and return its absolute path."""
    if not answer_sheets:
        raise ValueError("Cannot create a workbook without answer-sheet links.")

    output = output_path.expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output file must use the .xlsx extension.")
    if output.exists():
        raise FileExistsError(f"Output file already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MARKS_SHEET
    sheet.append(["Sheet No", "PDF Link"])

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for answer_sheet in answer_sheets:
        sheet.append([answer_sheet.sheet_no, answer_sheet.pdf_url])
        row = sheet.max_row
        sheet.cell(row=row, column=1).number_format = "@"
        link_cell = sheet.cell(row=row, column=2)
        link_cell.hyperlink = answer_sheet.pdf_url
        link_cell.style = "Hyperlink"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{sheet.max_row}"
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 90
    sheet.row_dimensions[1].height = 22
    _write_metadata(
        workbook,
        subject_code=subject_code,
        bundle_no=bundle_no,
        sheet_count=len(answer_sheets),
        exported_at=exported_at or datetime.now().astimezone(),
    )

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output.stem}.", suffix=".xlsx", dir=output.parent, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, output)
    except Exception:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()

    return output


def _metadata_values(workbook) -> dict[str, object]:
    if META_SHEET not in workbook.sheetnames:
        raise ValueError(
            "This is not a HITAM exam workbook. Run `hitam exam export` first."
        )
    sheet = workbook[META_SHEET]
    return {
        str(sheet.cell(row=row, column=1).value or "").strip(): sheet.cell(
            row=row, column=2
        ).value
        for row in range(1, sheet.max_row + 1)
    }


def _required_metadata(metadata: dict[str, object], label: str) -> str:
    value = str(metadata.get(label) or "").strip()
    if not value:
        raise ValueError(f"Workbook metadata is missing {label}.")
    return value


def _cell_decimal(cell: Cell, sheet_no: str, header: str) -> Decimal:
    if cell.data_type == "f":
        raise ValueError(f"Sheet {sheet_no}, {header}: formulas are not allowed.")
    value = cell.value
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"Sheet {sheet_no}, {header}: mark is blank.")
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError(  # noqa: TRY004 - invalid workbook content is a value error
            f"Sheet {sheet_no}, {header}: mark must be numeric."
        )
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"Sheet {sheet_no}, {header}: mark is invalid.") from exc
    if not result.is_finite() or result < 0:
        raise ValueError(
            f"Sheet {sheet_no}, {header}: mark must be a non-negative finite number."
        )
    return result


def _validate_pdf_url(value: object, sheet_no: str) -> str:
    url = str(value or "").strip()
    parsed = urlsplit(url)
    if parsed.scheme.lower() != "https" or not parsed.path.casefold().endswith(
        f"/{sheet_no}.pdf".casefold()
    ):
        raise ValueError(f"Sheet {sheet_no}: PDF Link is invalid or was changed.")
    return url


def read_marks_workbook(path: Path) -> MarksWorkbook:
    """Read and structurally validate a completed HITAM marks workbook."""
    source = path.expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Workbook not found: {source}")
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Marks workbook must use the .xlsx extension.")

    workbook = load_workbook(source, data_only=False, read_only=False)
    try:
        metadata = _metadata_values(workbook)
        version = _required_metadata(metadata, "HITAM Workbook Format")
        if version != WORKBOOK_FORMAT_VERSION:
            raise ValueError(
                "Unsupported HITAM workbook format. Run `hitam exam export` again."
            )
        subject_code = _required_metadata(metadata, "Subject Code").upper()
        bundle_no = _required_metadata(metadata, "Bundle No")
        try:
            expected_count = int(metadata.get("Expected Sheet Count"))
        except (TypeError, ValueError) as exc:
            raise ValueError("Workbook metadata has an invalid sheet count.") from exc

        if MARKS_SHEET not in workbook.sheetnames:
            raise ValueError(f"Workbook is missing the {MARKS_SHEET!r} sheet.")
        sheet = workbook[MARKS_SHEET]
        if sheet.cell(1, 1).value != "Sheet No" or sheet.cell(1, 2).value != "PDF Link":
            raise ValueError("The first two headers must be Sheet No and PDF Link.")

        last_header_column = max(
            (
                column
                for column in range(3, sheet.max_column + 1)
                if sheet.cell(1, column).value not in (None, "")
            ),
            default=2,
        )
        if last_header_column == 2:
            raise ValueError(
                "Add question columns such as Q1, Q2, and Q3 before upload."
            )

        question_headers: list[str] = []
        question_keys: list[str] = []
        for column in range(3, last_header_column + 1):
            header = str(sheet.cell(1, column).value or "").strip()
            if not header:
                raise ValueError("Question headers cannot contain gaps.")
            key = normalise_question_label(header)
            if not key:
                raise ValueError(f"Invalid question header: {header!r}.")
            if key in question_keys:
                raise ValueError(f"Duplicate question header: {header!r}.")
            question_headers.append(header)
            question_keys.append(key)

        students: list[StudentMarks] = []
        seen: set[str] = set()
        for row in range(2, sheet.max_row + 1):
            raw_sheet_no = sheet.cell(row, 1).value
            raw_pdf_url = sheet.cell(row, 2).value
            if raw_sheet_no in (None, "") and raw_pdf_url in (None, ""):
                continue
            if (
                sheet.cell(row, 1).data_type == "f"
                or sheet.cell(row, 2).data_type == "f"
            ):
                raise ValueError(
                    f"Row {row}: formulas are not allowed in identity columns."
                )
            sheet_no = str(raw_sheet_no or "").strip()
            if not sheet_no:
                raise ValueError(f"Row {row}: Sheet No is blank.")
            if sheet_no in seen:
                raise ValueError(f"Duplicate Sheet No: {sheet_no}.")
            seen.add(sheet_no)
            pdf_url = _validate_pdf_url(raw_pdf_url, sheet_no)
            marks = {
                key: _cell_decimal(sheet.cell(row, column), sheet_no, header)
                for column, (header, key) in enumerate(
                    zip(question_headers, question_keys, strict=True), start=3
                )
            }
            students.append(StudentMarks(sheet_no, pdf_url, marks))

        if len(students) != expected_count:
            raise ValueError(
                f"Workbook contains {len(students)} students; expected {expected_count}."
            )
        return MarksWorkbook(
            path=source,
            subject_code=subject_code,
            bundle_no=bundle_no,
            expected_sheet_count=expected_count,
            question_headers=tuple(question_headers),
            students=tuple(students),
        )
    finally:
        workbook.close()
