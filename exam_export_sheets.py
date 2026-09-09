from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pathlib import Path
from datetime import datetime
import re

PORTAL_URL = "https://hitamexams.org/MidEvaluation/"

# -----------------------------
# CONFIGURATION
# -----------------------------
USER_CODE = "HTM5003"
SUBJECT_CODE = "24PC5CD04"
BUNDLE_NO = "FSD-5-5-MID-I-AUG26-3"
BUNDLE_KEY = "4324E2106101412"

# The generated file is directly compatible with the marks-upload script.
OUTPUT_FILE = "marks.xlsx"
QUESTION_COUNT = 11


# -----------------------------
# LOGIN
# -----------------------------
def login(page):
    page.goto(
        PORTAL_URL,
        wait_until="domcontentloaded",
        timeout=30000,
    )

    page.get_by_role("button", name="FACULTY LOGIN").click()

    page.locator("select").select_option(label="B.Tech")
    page.get_by_role("button", name="VALIDATE").click()

    page.get_by_label("USER CODE").fill(USER_CODE)
    page.get_by_role("button", name="VALIDATE").click()

    page.get_by_label("BUNDLE NO").fill(BUNDLE_NO)
    page.get_by_role("button", name="VALIDATE").click()

    page.get_by_label("BUNDLE KEY").fill(BUNDLE_KEY)
    page.get_by_role("button", name="VALIDATE").click()

    print("Login completed.")


# -----------------------------
# GET SHEET NUMBERS
# -----------------------------
def getSheetNumbers(page):
    sheet_buttons = page.locator(
        "button.MuiButton-outlined"
    ).filter(
        has_text=re.compile(r"^\s*\d{6}\s*$")
    )

    sheet_buttons.first.wait_for(
        state="visible",
        timeout=30000,
    )

    sheet_numbers = []
    seen = set()

    for text in sheet_buttons.all_inner_texts():
        text = text.strip()

        if re.fullmatch(r"\d{6}", text) and text not in seen:
            seen.add(text)
            sheet_numbers.append(text)

    print("Total sheets found:", len(sheet_numbers))

    return sheet_numbers


# -----------------------------
# CAPTURE PDF LINK
# -----------------------------
def getPdfLink(page, sheet_no):
    button = page.locator(
        "button.MuiButton-outlined"
    ).filter(
        has_text=re.compile(
            rf"^\s*{re.escape(sheet_no)}\s*$"
        )
    ).first

    with page.expect_response(
        lambda response: (
            response.url.lower().endswith(
                f"/{sheet_no}.pdf".lower()
            )
            and response.status == 200
        ),
        timeout=30000,
    ) as response_info:
        button.click()

    pdf_response = response_info.value

    # Make sure the PDF response body has been completely received.
    pdf_response.body()

    return pdf_response.url


# -----------------------------
# COLLECT ALL SHEETS
# -----------------------------
def collectAnswerSheets(page, sheet_numbers):
    answer_sheets = []

    for index, sheet_no in enumerate(sheet_numbers, start=1):
        print(
            f"[{index}/{len(sheet_numbers)}] "
            f"Opening sheet {sheet_no}..."
        )

        pdf_link = getPdfLink(
            page,
            sheet_no,
        )

        answer_sheets.append({
            "sheet_no": sheet_no,
            "pdf_link": pdf_link,
        })

        print("PDF:", pdf_link)

    return answer_sheets


# -----------------------------
# CREATE EXCEL WORKBOOK
# -----------------------------
def createExcel(answer_sheets, output_file):
    output_path = Path(output_file).resolve()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Answer sheets"

    headers = [
        "Sheet No",
        "PDF Link",
        *[
            f"Q{i}"
            for i in range(1, QUESTION_COUNT + 1)
        ],
    ]

    sheet.append(headers)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for answer_sheet in answer_sheets:
        sheet.append([
            answer_sheet["sheet_no"],
            answer_sheet["pdf_link"],
            *([None] * QUESTION_COUNT),
        ])

        row = sheet.max_row

        # Keep sheet number as text.
        sheet.cell(
            row=row,
            column=1,
        ).number_format = "@"

        # Make PDF URL clickable.
        pdf_cell = sheet.cell(
            row=row,
            column=2,
        )
        pdf_cell.hyperlink = answer_sheet["pdf_link"]
        pdf_cell.style = "Hyperlink"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 90

    for column in range(3, 3 + QUESTION_COUNT):
        sheet.column_dimensions[
            sheet.cell(1, column).column_letter
        ].width = 10

    sheet.row_dimensions[1].height = 22

    # A small hidden metadata sheet is useful for checking the workbook later.
    metadata = workbook.create_sheet("__hitam_meta")
    metadata.append(["Subject Code", SUBJECT_CODE])
    metadata.append(["Bundle No", BUNDLE_NO])
    metadata.append(["Expected Sheet Count", len(answer_sheets)])
    metadata.append([
        "Exported At",
        datetime.now().astimezone().isoformat(
            timespec="seconds"
        ),
    ])
    metadata.sheet_state = "veryHidden"

    workbook.save(output_path)
    workbook.close()

    print()
    print("Excel file created successfully:")
    print(output_path)

    return output_path


# -----------------------------
# MAIN
# -----------------------------
if __name__ == "__main__":
    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=False
        )

        page = browser.new_page()

        try:
            login(page)

            sheet_numbers = getSheetNumbers(page)

            answer_sheets = collectAnswerSheets(
                page,
                sheet_numbers,
            )

            createExcel(
                answer_sheets,
                OUTPUT_FILE,
            )

        finally:
            browser.close()
