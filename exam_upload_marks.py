from playwright.sync_api import sync_playwright
import re
from openpyxl import load_workbook

PORTAL_URL = "https://hitamexams.org/MidEvaluation/"

USER_CODE = "HTM5003"
BUNDLE_NO = "FSD-5-5-MID-I-AUG26-3"
BUNDLE_KEY = "4324E2106101412"
EXCEL_FILE_PATH = "upload_marks.xlsx"

def login(page):
    page.goto(PORTAL_URL)

    # Faculty Login
    page.get_by_role("button", name="FACULTY LOGIN").click()

    # Select program
    page.locator("select").select_option(label="B.Tech")

    page.get_by_role("button", name="VALIDATE").click()

    # User Code
    page.get_by_label("USER CODE").fill(USER_CODE)
    page.get_by_role("button", name="VALIDATE").click()

    # Bundle Number
    page.get_by_label("BUNDLE NO").fill(BUNDLE_NO)
    page.get_by_role("button", name="VALIDATE").click()

    # Bundle Key
    page.get_by_label("BUNDLE KEY").fill(BUNDLE_KEY)
    page.get_by_role("button", name="VALIDATE").click()

    print("Login completed.")
def getSheetNumbers(page):
    sheet_buttons = page.locator(
        "button.MuiButton-outlined"
    ).filter(
        has_text=re.compile(r"^\s*\d{6}\s*$")
    )

    sheet_buttons.first.wait_for(
        state="visible",
        timeout=30000,
    )

    sheet_numbers = [
        text.strip()
        for text in sheet_buttons.all_inner_texts()
        if re.fullmatch(r"\d{6}", text.strip())
    ]

    print("Total sheets found:", len(sheet_numbers))

    return sheet_numbers
def waitForPdf(page, sheet_no):
    with page.expect_response(
        lambda response: (
            response.url.lower().endswith(f"/{sheet_no}.pdf")
            and response.status == 200
        ),
        timeout=30000,
    ) as pdf_response_info:

        button = page.locator(
            "button.MuiButton-outlined",
            has_text=sheet_no,
        ).first

        button.click()

    pdf_response = pdf_response_info.value

    # Wait until the complete PDF body has been received
    pdf_response.body()

    return pdf_response
def fillMarksAndSave(page, student):
    sheet_no = student["sheet_no"]
    marks = student["marks"]

    marks_inputs = page.locator(
        "tbody input[type='text']"
    )

    marks_inputs.first.wait_for(
        state="visible",
        timeout=30000,
    )

    count = marks_inputs.count()

    print(
        f"{sheet_no}: Marks fields found: {count}"
    )

    if count != len(marks):
        raise Exception(
            f"{sheet_no}: portal has {count} fields "
            f"but Excel has {len(marks)} marks"
        )

    # Clear
    for i in range(count):
        marks_inputs.nth(i).fill("")

    print(f"{sheet_no}: All marks cleared")

    page.wait_for_timeout(500)

    # Fill
    expected_values = []

    for i, mark in enumerate(marks):
        if mark is None:
            value = ""
        elif isinstance(mark, float) and mark.is_integer():
            value = str(int(mark))
        else:
            value = str(mark)

        expected_values.append(value)

        marks_inputs.nth(i).fill(value)

        print(
            f"{sheet_no} Q{i + 1}: {value}"
        )

    # Verify
    for i, expected in enumerate(expected_values):
        actual = marks_inputs.nth(i).input_value()

        if actual != expected:
            raise Exception(
                f"{sheet_no} Q{i + 1}: "
                f"expected {expected}, found {actual}"
            )

    print(f"{sheet_no}: All marks verified")

    # Save and WAIT for server
    with page.expect_response(
        lambda response: (
            "MIDSecuritySavingMarksNew"
            in response.url
        ),
        timeout=30000,
    ) as save_info:

        page.get_by_role(
            "button",
            name="Save",
            exact=True,
        ).click()

    save_response = save_info.value

    print(
        f"{sheet_no}: Save response "
        f"{save_response.status}"
    )

    if not save_response.ok:
        raise Exception(
            f"{sheet_no}: save failed with "
            f"HTTP {save_response.status}"
        )

    print(f"{sheet_no}: Marks saved successfully")
def loadMarksFromExcel(file_path):
    workbook = load_workbook(
        file_path,
        data_only=True
    )

    sheet = workbook["Answer sheets"]

    students = {}

    # Skip header row
    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        sheet_no = row[0]
        pdf_link = row[1]

        if sheet_no is None:
            continue

        sheet_no = str(sheet_no).strip()

        # Q1 to Q11
        marks = list(row[2:13])

        students[sheet_no] = {
            "sheet_no": sheet_no,
            "pdf_link": str(pdf_link).strip() if pdf_link else "",
            "marks": marks,
        }

    workbook.close()

    print(
        "Students loaded from Excel:",
        len(students)
    )

    return students

if __name__ == "__main__":
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        page = browser.new_page()
        # Load marks from Excel before starting the automation
        students = loadMarksFromExcel(EXCEL_FILE_PATH)
        # Login on the portal
        login(page)
        # Get all sheet numbers after login
        sheet_numbers = getSheetNumbers(page)
        # Open each sheet and wait for its PDF to load.
        for sheet_no in sheet_numbers:
            print("\n----------------------------")
            print("Opening sheet:", sheet_no)
            student = students.get(sheet_no)
            if student is None:
                print(f"No Excel data found for {sheet_no}")
                continue
            pdf_response = waitForPdf(
                page,
                sheet_no
            )
            print("PDF loaded:",sheet_no,pdf_response.url)
            fillMarksAndSave(
                page,
                student
            )

        # Keep browser open
        # input("Press Enter to close browser...")
        browser.close()
        

